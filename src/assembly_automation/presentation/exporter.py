"""Deterministic editable PowerPoint export for assembly and monopart artifacts."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import json
from pathlib import Path
import re
from typing import Any, Mapping, Sequence

from assembly_automation.workflows.definitions.app_v3 import WorkflowPaths


@dataclass(frozen=True)
class EngineeringPowerPointContext:
    """Active saved artifacts and their preprocessing image directory."""

    session_root: Path
    assembly: Mapping[str, Any]
    bom: Mapping[str, Any]
    images_root: Path
    assembly_revision: str
    monoparts_revision: str

    @classmethod
    def from_session(cls, session_root: str | Path) -> "EngineeringPowerPointContext":
        paths = WorkflowPaths(Path(session_root).resolve())
        assembly_revision = paths.active_revision("assembly")
        monoparts_revision = paths.active_revision("monoparts")
        assembly = _load_object(paths.assembly(assembly_revision) / "assembly_overview.json",
                                "active assembly overview")
        bom = _load_object(paths.monoparts(monoparts_revision) / "bom.json", "active BOM")
        return cls(Path(session_root).resolve(), assembly, bom,
                   paths.preprocessing / "images", assembly_revision, monoparts_revision)


def build_report(context: EngineeringPowerPointContext | str | Path,
                 sections: Sequence[str] = ("assembly", "monopart"),
                 part_ids: Sequence[str] | None = None) -> Path:
    """Build and save one versioned report from one active artifact snapshot."""
    _require_pptx()
    if not isinstance(context, EngineeringPowerPointContext):
        context = EngineeringPowerPointContext.from_session(context)
    selected_sections = tuple(sections)
    if not selected_sections or any(item not in {"assembly", "monopart"} for item in selected_sections):
        raise ValueError("sections must contain only 'assembly' and/or 'monopart'")
    parts = _parts(context.bom, part_ids)
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    _remove_default_slides(presentation)
    _cover(presentation, context, len(parts))
    if "assembly" in selected_sections:
        _render_assembly(presentation, context)
    if "monopart" in selected_sections:
        for part in parts:
            _render_part(presentation, context, part)
    output_dir = context.session_root / "09_user_agent" / "powerpoint_exports"
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    output = output_dir / f"engineering_report_{stamp}.pptx"
    presentation.save(output)
    return output


def _render_assembly(presentation: Any, context: EngineeringPowerPointContext) -> None:
    slide = _slide(presentation, "Assembly")
    assembly_name = _display(context.assembly.get("assembly_name_guess"), "Assembly")
    _title(slide, f"{assembly_name} - Overview")
    _add_image_or_placeholder(slide, _find_image(context.images_root, "assembly", "exploded"),
                              "Assembly exploded image unavailable",
                              left=8.65, top=1.15, width=4.0, height=2.953)
    _add_image_or_placeholder(slide, _find_image(context.images_root, "assembly", "iso1",
                                                 exclude=("exploded",)),
                              "Assembly iso1 image unavailable",
                              left=8.65, top=4.25, width=4.0, height=2.953)
    _add_card(slide, "Assembly architecture", _lines(context.assembly.get("assembly_description")),
              left=0.65, top=1.15, width=3.6, height=1.65, accent="green", bullets=False)
    _add_card(slide, "Primary function", _lines(context.assembly.get("primary_function"), limit=3),
              left=4.5, top=1.15, width=3.6, height=1.65, accent="blue")
    interfaces = _interface_lines(context.assembly.get("interfaces"))
    if interfaces:
        _add_card(slide, "Key mechanical interfaces", interfaces[:3],
                  left=0.65, top=3.0, width=3.6, height=3.25, accent="amber")
    uncertainties = _lines(context.assembly.get("uncertainties"), limit=2)
    if uncertainties:
        _add_card(slide, "Review focus", uncertainties,
                  left=4.5, top=3.0, width=3.6, height=3.25, accent="red")

    components = context.assembly.get("partslist")
    if isinstance(components, list) and any(isinstance(item, Mapping) for item in components):
        _component_slide(presentation, context, components)


def _render_part(presentation: Any, context: EngineeringPowerPointContext,
                 part: Mapping[str, Any]) -> None:
    slide = _slide(presentation, "Monopart")
    analysis = part.get("part_analysis") if isinstance(part.get("part_analysis"), Mapping) else {}
    name = _display(analysis.get("part_name_guess") or part.get("name"), "Unnamed part")
    _title(slide, f"{name} · {part.get('part_id', 'unknown ID')}")
    part_id = str(part.get("part_id", ""))
    _add_image_or_placeholder(slide, _find_image(context.images_root, part_id, "iso1"),
                              "Part iso1 image unavailable",
                              left=8.65, top=1.15, width=4.0, height=2.95)
    _add_image_or_placeholder(slide, _find_image(context.images_root, part_id, "highlighted"),
                              "Part highlighted image unavailable",
                              left=8.65, top=4.15, width=4.0, height=2.95)
    _add_card(slide, "Role in the assembly", _lines(analysis.get("part_identification")),
              left=.65, top=1.15, width=3.6, height=2.025, accent="green", bullets=False)
    _add_card(slide, "Engineering readout", _lines(analysis.get("intrinsic_summary"), limit=3),
              left=4.5, top=1.15, width=3.6, height=2.025, accent="blue")
    _add_card(slide, "Decisive geometry", _lines(analysis.get("geometric_characteristics"), limit=4),
              left=.65, top=3.4, width=3.6, height=2.9625, accent="green")
    _add_card(slide, "Handling and gripping",
              _lines(analysis.get("handling_implications"), limit=2)
              + _lines(analysis.get("gripping_analysis"), limit=2),
              left=4.5, top=3.4, width=3.6, height=2.9625, accent="blue")


def _component_slide(presentation: Any, context: EngineeringPowerPointContext,
                     components: Sequence[Mapping[str, Any]]) -> None:
    slide = _slide(presentation, "Components")
    assembly_name = _display(context.assembly.get("assembly_name_guess"), "Assembly")
    _title(slide, f"{assembly_name} - Bill of Material")
    from pptx.util import Inches, Pt
    table_left, table_top = .65, 1.35
    table = slide.shapes.add_table(max(1, len(components) + 1), 5,
                                   Inches(table_left), Inches(table_top),
                                   Inches(12.0), Inches(5.15)).table
    widths = (1.05, 2.05, 2.0, .75, 6.15)
    for column, width in enumerate(widths):
        table.columns[column].width = Inches(width)
    for index, header in enumerate(("ISO", "Instance IDs", "Component", "Qty.", "Assembly role")):
        _cell(table.cell(0, index), header, bold=True)
    table.rows[0].height = Inches(.42)
    for row, component in enumerate(components, 1):
        instance_ids = component.get("instance_ids")
        values = ("",
                  ", ".join(str(item) for item in instance_ids) if isinstance(instance_ids, list) else "—",
                  component.get("name"), len(instance_ids) if isinstance(instance_ids, list) else "—",
                  " ".join(_lines(component.get("assembly_role"), limit=2)))
        for column, value in enumerate(values):
            if column == 0:
                table.cell(row, column).text = ""
            else:
                _cell(table.cell(row, column), value, bold=False)
        table.rows[row].height = Inches(.7)
        identifier = str((instance_ids or [""])[0]) if isinstance(instance_ids, list) else ""
        image = _find_image(context.images_root, identifier or str(component.get("name", "")), "iso1")
        if image:
            slide.shapes.add_picture(str(image), Inches(table_left + .17),
                                     Inches(table_top + .42 + (row - 1) * .7 + .06),
                                     height=Inches(.58))
    _style_table(table)


def _lines(value: Any, *, limit: int | None = None) -> list[str]:
    """Normalize concise schema fields without exposing nested JSON to readers."""
    if value is None:
        return []
    values = value if isinstance(value, list) else [value]
    lines = [_strip_list_marker(str(item).strip()) for item in values
             if isinstance(item, (str, int, float)) and str(item).strip()]
    return lines[:limit] if limit is not None else lines


def _interface_lines(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    lines: list[str] = []
    for item in value:
        if isinstance(item, Mapping):
            lines.extend(_lines(item.get("statements"), limit=2))
    return lines


def _add_card(slide: Any, title: str, lines: Sequence[str], *, left: float, top: float,
              width: float, height: float, accent: str, bullets: bool = True) -> None:
    """Add a restrained editable engineering-information card."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    palette = {
        "green": RGBColor(24, 75, 68),
        "blue": RGBColor(53, 105, 155),
        "amber": RGBColor(190, 126, 30),
        "red": RGBColor(170, 72, 60),
    }
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(220, 227, 225)
    shape.line.width = Pt(.75)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(.07), Inches(height))
    bar.fill.solid()
    bar.fill.fore_color.rgb = palette[accent]
    bar.line.fill.background()
    heading = slide.shapes.add_textbox(Inches(left + .22), Inches(top + .12),
                                       Inches(width - .35), Inches(.24))
    heading_paragraph = heading.text_frame.paragraphs[0]
    heading_paragraph.text = title.upper()
    heading_paragraph.font.name = "Aptos"
    heading_paragraph.font.size = Pt(9)
    heading_paragraph.font.bold = True
    heading_paragraph.font.color.rgb = palette[accent]
    if not lines:
        return
    body = slide.shapes.add_textbox(Inches(left + .22), Inches(top + .42),
                                    Inches(width - .4), Inches(max(height - .52, .1)))
    frame = body.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _add_markdown_paragraph(paragraph, line, is_bullet=bullets)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(11)
        paragraph.space_after = Pt(4)


def _style_table(table: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Pt
    for row_index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            cell.margin_left = cell.margin_right = Pt(4)
            cell.margin_top = cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = (RGBColor(53, 105, 155) if row_index == 0
                                        else RGBColor(255, 255, 255))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.LEFT if column == 4 else PP_ALIGN.CENTER
                paragraph.font.size = Pt(9)
                paragraph.font.name = "Aptos"
                paragraph.font.color.rgb = (RGBColor(255, 255, 255) if row_index == 0
                                            else RGBColor(34, 42, 48))
            tc_pr = cell._tc.get_or_add_tcPr()
            for edge_name in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                edge = next((child for child in tc_pr
                             if child.tag.rsplit("}", 1)[-1] == edge_name.rsplit(":", 1)[-1]),
                            None)
                if edge is None:
                    edge = OxmlElement(edge_name)
                    tc_pr.append(edge)
                edge.set("w", "12700")
                solid = OxmlElement("a:solidFill")
                color = OxmlElement("a:srgbClr")
                color.set("val", "000000")
                solid.append(color)
                edge.append(solid)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def _add_json_text(slide: Any, value: Any, *, left: float, top: float,
                   width: float, height: float) -> None:
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    for index, (line, is_bullet) in enumerate(_flatten(value)):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _add_markdown_paragraph(paragraph, line, is_bullet=is_bullet)
        paragraph.font.size = Pt(14 if index == 0 else 11)
        paragraph.font.name = "Aptos"


def _flatten(value: Any, prefix: str = "", is_bullet: bool = False) -> list[tuple[str, bool]]:
    if isinstance(value, Mapping):
        lines: list[tuple[str, bool]] = []
        for key, item in value.items():
            label = f"{prefix}{str(key).replace('_', ' ')}:"
            if isinstance(item, list):
                lines.append((label, False))
                lines.extend(_flatten(item, is_bullet=True))
            elif isinstance(item, Mapping):
                lines.append((label, False))
                lines.extend(_flatten(item))
            else:
                lines.extend(_flatten(item, f"{label} "))
        return lines
    if isinstance(value, list):
        lines: list[tuple[str, bool]] = []
        for item in value:
            item_lines = str(item).splitlines() or [""]
            for line in item_lines:
                lines.append((f"{prefix}{_strip_list_marker(line.strip())}", True))
        return lines or [(f"{prefix}—", is_bullet)]
    return [(f"{prefix}{_display(value)}", is_bullet)]


def _strip_list_marker(value: str) -> str:
    return re.sub(r"^(?:[-*]|\d+[.)])\s+", "", value)


def _add_markdown_paragraph(paragraph: Any, value: str, *, is_bullet: bool) -> None:
    """Write Markdown inline emphasis as editable runs and list items as bullets."""
    if is_bullet:
        from pptx.oxml.xmlchemy import OxmlElement

        properties = paragraph._p.get_or_add_pPr()
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        properties.append(bullet)
    paragraph.clear()
    token_pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    position = 0
    for match in token_pattern.finditer(value):
        if match.start() > position:
            paragraph.add_run().text = value[position:match.start()]
        token = match.group(0)
        run = paragraph.add_run()
        if token.startswith("**"):
            run.text = token[2:-2]
            run.font.bold = True
        else:
            run.text = token[1:-1]
            run.font.italic = True
        position = match.end()
    if position < len(value):
        paragraph.add_run().text = value[position:]


def _find_image(root: Path, identifier: str, view: str,
                exclude: Sequence[str] = ()) -> Path | None:
    if not root.is_dir():
        return None
    token = identifier.lower().replace(" ", "_")
    candidates = sorted(path for path in root.rglob("*")
                        if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg"}
                        and view.lower() in path.stem.lower()
                        and not any(token.lower() in path.stem.lower() for token in exclude))
    exact = [path for path in candidates
             if token and (token in path.stem.lower()
                           or token == path.parent.name.lower())]
    return (exact or candidates)[0] if (exact or candidates) else None


def _parts(bom: Mapping[str, Any], part_ids: Sequence[str] | None) -> list[Mapping[str, Any]]:
    parts = [part for part in bom.get("parts", []) if isinstance(part, Mapping)]
    if part_ids is None:
        return parts
    requested = set(part_ids)
    unknown = requested - {str(part.get("part_id")) for part in parts}
    if unknown:
        raise ValueError(f"Unknown part IDs: {sorted(unknown)}")
    return [part for part in parts if str(part.get("part_id")) in requested]


def _slide(presentation: Any, section: str) -> Any:
    from pptx.dml.color import RGBColor
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(247, 250, 249)
    _add_logo(slide)
    return slide


def _title(slide: Any, text: str) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.45), Inches(12), Inches(0.6))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos Display"
    paragraph.font.size = Pt(27)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor(24, 75, 68)


def _placeholder(slide: Any, text: str) -> None:
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(8.55), Inches(3.0), Inches(4.2), Inches(1.0))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(16)


def _cover(presentation: Any, context: EngineeringPowerPointContext, count: int) -> None:
    slide = _slide(presentation, "Cover")
    assembly_name = _display(context.assembly.get("assembly_name_guess"), "Assembly")
    _title(slide, f"{assembly_name} - Engineering Report")
    image = _find_image(context.images_root, "assembly", "iso1", exclude=("exploded",))
    if image:
        from pptx.util import Inches
        slide.shapes.add_picture(str(image), Inches(3.65), Inches(1.35),
                                 width=Inches(6.0))
    else:
        _placeholder(slide, "Assembly iso1 image unavailable")


def _add_logo(slide: Any) -> None:
    logo = Path(__file__).resolve().parents[3] / "logos" / "HB.png"
    if logo.is_file():
        from pptx.util import Inches
        slide.shapes.add_picture(str(logo), Inches(11.75), Inches(.16), width=Inches(1.35))


def _add_image_or_placeholder(slide: Any, image: Path | None, text: str, *,
                              left: float, top: float, width: float,
                              height: float | None = None) -> None:
    from pptx.util import Inches
    if image:
        kwargs = {"width": Inches(width)}
        if height is not None:
            kwargs = {"height": Inches(height)}
        slide.shapes.add_picture(str(image), Inches(left), Inches(top), **kwargs)
    else:
        _placeholder(slide, text)


def _cell(cell: Any, value: Any, *, bold: bool) -> None:
    cell.text = _display(value)
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Aptos"
            run.font.bold = bold


def _display(value: Any, fallback: str = "—") -> str:
    if value is None or value == "":
        return fallback
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, indent=2)
    return str(value)


def _load_object(path: Path, label: str) -> dict[str, Any]:
    if not path.is_file():
        raise FileNotFoundError(f"{label} is not available: {path}")
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{label} must be a JSON object")
    return value


def _remove_default_slides(presentation: Any) -> None:
    for slide in list(presentation.slides):
        presentation.slides._sldIdLst.remove(slide._element)


def _require_pptx() -> None:
    try:
        import pptx  # noqa: F401
    except ImportError as exc:
        raise RuntimeError("PowerPoint export requires python-pptx") from exc
