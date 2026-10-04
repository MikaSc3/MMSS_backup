import json
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from assembly_automation.presentation import build_report
from assembly_automation.workflows.definitions.app_v3 import WorkflowPaths


class PowerPointExportTests(unittest.TestCase):
    def test_exports_active_assembly_and_all_distinct_parts(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths = WorkflowPaths(root)
            (root / "artifact_registry.json").write_text(json.dumps(
                {"active": {"assembly": "r003", "monoparts": "r004"}}), encoding="utf-8")
            assembly = paths.assembly("r003") / "assembly_overview.json"
            assembly.parent.mkdir(parents=True)
            assembly.write_text(json.dumps({
                "assembly_name_guess": "Test fixture",
                "assembly_description": "Two parts",
                "primary_function": ["**Locate** parts", "Maintain alignment"],
                "partslist": ["Base and pin"],
                "extra": {"unit": "mm"},
            }), encoding="utf-8")
            bom = paths.monoparts("r004") / "bom.json"
            bom.parent.mkdir(parents=True)
            bom.write_text(json.dumps({"parts": [
                {"part_id": "part_001", "name": "Base", "quantity": 1,
                 "geometry": {"length": 20},
                 "part_analysis": {"part_name_guess": "Base",
                                   "part_identification": "A **flat** base",
                                   "intrinsic_summary": ["**Geometry:** Flat", "*Material:* Steel"]}},
                {"part_id": "part_002", "name": "Pin", "quantity": 2,
                 "part_analysis": {"part_name_guess": "Pin",
                                   "gripping_analysis": "Side grip",
                                   "part_identification": "Round pin",
                                   "intrinsic_summary": ["Cylindrical", "Side grip"]}},
            ], "instances": []}), encoding="utf-8")

            output = build_report(root)

            self.assertTrue(output.is_file())
            self.assertEqual(output.parent.name, "powerpoint_exports")
            from pptx import Presentation
            presentation = Presentation(output)
            self.assertEqual(len(presentation.slides), 4)
            text = "\n".join(
                shape.text for slide in presentation.slides for shape in slide.shapes
                if hasattr(shape, "text")
            )
            self.assertIn("Test fixture", text)
            self.assertIn("part_001", text)
            self.assertIn("part_002", text)
            self.assertIn("Two parts", text)
            self.assertIn("Locate parts", text)
            self.assertIn("Maintain alignment", text)
            self.assertNotIn("Base and pin", text)
            self.assertNotIn("length", text)

            assembly_slide = presentation.slides[1]
            body = next(shape for shape in assembly_slide.shapes
                        if hasattr(shape, "text_frame") and "Locate" in shape.text)
            bullet_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar"
            self.assertTrue(any(
                paragraph._p.pPr is not None and paragraph._p.pPr.find(bullet_tag) is not None
                for paragraph in body.text_frame.paragraphs
            ))
            locate_run = next(
                run for paragraph in body.text_frame.paragraphs
                for run in paragraph.runs if run.text == "Locate"
            )
            self.assertTrue(locate_run.font.bold)
